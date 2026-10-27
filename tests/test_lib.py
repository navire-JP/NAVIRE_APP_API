"""
NAVIRE Lib contre une base SQLite : import d'un cours .docx par la console,
lecture par un élève (contenu flouté au point de passage), questions sans les
réponses, correction, gain d'Elo, verrou d'Elo du point suivant.

Lancer : pytest tests/test_lib.py
"""

import io
import os
import uuid

import pytest

for k, v in {
    "OPENAI_API_KEY": "x", "BREVO_API_KEY_MEOLES": "x",
    "STRIPE_SECRET_KEY_MEOLES": "x", "STRIPE_WEBHOOK_SECRET_MEOLES": "x",
    "STRIPE_SECRET_KEY": "sk_test_x", "STRIPE_WEBHOOK_SECRET": "whsec_x",
    "DATABASE_URL": "sqlite:///./test_prepa_adjuris.db",
}.items():
    os.environ.setdefault(k, v)

from fastapi.testclient import TestClient  # noqa: E402

ADMIN = {"X-Admin-Code": "THORKISHERE"}

NOTIONS = [
    ("Distribution directe", "Succursales, salariés, vente à distance. Maîtrise totale."),
    ("Distribution indirecte", "Intermédiaires indépendants, pour le compte du producteur ou pour leur compte."),
    ("Réseau", "Ensemble de contrats conclus selon un modèle uniforme."),
    ("Loi Doubin", "Art. L. 330-3 C. com. Document d'information précontractuelle."),
    ("Franchise", "Signes distinctifs, savoir-faire, assistance."),
    ("Concession", "Exclusivité territoriale contre approvisionnement exclusif."),
]


def _cours_docx() -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = docx.Document()

    def p(text, bold=False, italic=False, center=False):
        par = d.add_paragraph()
        r = par.add_run(text)
        r.bold, r.italic = bold, italic
        if center:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return par

    def lead(n, titre, corps):
        par = d.add_paragraph()
        par.add_run(f"{n}. {titre}.- ").bold = True
        par.add_run(corps)

    p("NAVIRE LIB · MASTER 1", center=True)
    p("DROIT DE LA DISTRIBUTION", bold=True, center=True)
    p("Contrats de distribution et concurrence", center=True)
    p("À jour au 24 octobre 2026", italic=True, center=True)
    p("PLAN SOMMAIRE", bold=True)
    p("ÉTAPE 1. NOTIONS\t3")
    p("ÉTAPE 2. CONTRATS\t5")
    texte = "Le producteur peut distribuer lui-même ses produits selon l'article L. 330-3 du Code de commerce. " * 6
    for e, (titre, notions) in enumerate([("NOTIONS", NOTIONS[:3]), ("CONTRATS", NOTIONS[3:])], 1):
        p(f"ÉTAPE {e}. {titre}", bold=True)
        p("Objectif de l'étape : comprendre.", italic=True)
        p("TITRE 1. PREMIER TITRE", bold=True)
        for i in range(3):
            lead(e * 10 + i, f"Paragraphe {e}-{i}", texte)
        p("POINTS DE PASSAGE", bold=True)
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Notion", "Ce qu'il faut retenir"
        for n, r in notions:
            row = t.add_row()
            row.cells[0].text, row.cells[1].text = n, r
    p("SOURCES", bold=True)
    p("Code de commerce, art. L. 330-3.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    import app.main as main
    import app.routers.lib as lib
    from app.db.database import Base, engine

    lib.LIB_DIR = tmp_path_factory.mktemp("lib")
    engine.dispose()
    Base.metadata.create_all(bind=engine)
    from app.db.lib_migration import run_lib_migration
    run_lib_migration()  # base de test d'une version précédente
    # Sans `with` : le lifespan (pgvector, migrations Postgres) ne tourne pas.
    yield TestClient(main.app)
    engine.dispose()


def _user(username, plan="membre"):
    from app.db.database import SessionLocal
    from app.db.models import User
    from app.core.security import create_access_token

    db = SessionLocal()
    u = User(email=f"{username}-{uuid.uuid4().hex[:8]}@test.fr", username=f"{username}{uuid.uuid4().hex[:4]}", plan=plan)
    db.add(u)
    db.commit()
    db.refresh(u)
    h = {"Authorization": f"Bearer {create_access_token(str(u.id))}"}
    db.close()
    return u.id, h


def _reponse(client, h, cp_id, juste=True):
    """La question de l'essai et sa bonne (ou mauvaise) réponse, lue en base."""
    from app.db.database import SessionLocal
    from app.db.models import LibCheckpoint

    r = client.get(f"/lib/checkpoints/{cp_id}/question", headers=h)
    assert r.status_code == 200, r.text
    q = r.json()["question"]
    assert "answer" not in q and "accepted" not in q
    db = SessionLocal()
    full = next(x for x in db.get(LibCheckpoint, cp_id).questions if x["id"] == q["id"])
    db.close()
    if full["type"] == "qcm":
        a = full["answer"] if juste else (full["answer"] + 1) % len(full["choices"])
    else:
        a = full["accepted"][0].lower() if juste else "aucune idée"
    return {"question_id": q["id"], "answer": a}


def _repondre(client, h, cp_id, juste=True):
    return client.post(f"/lib/checkpoints/{cp_id}/submit", headers=h, json=_reponse(client, h, cp_id, juste)).json()


def test_parse_docx_structure():
    from app.services.lib_parser import parse_document, detect_checkpoints

    r = parse_document(_cours_docx(), "cours.docx")
    assert r["meta"]["titre"] == "DROIT DE LA DISTRIBUTION"
    assert r["meta"]["maj"].startswith("À jour")
    types = [b["t"] for b in r["blocks"]]
    assert types.count("h1") == 3            # deux étapes + sources
    assert types.count("pdp") == 2 and types.count("table") == 2
    assert not any("PLAN SOMMAIRE" in b["text"] for b in r["blocks"])  # sommaire écarté
    p = next(b for b in r["blocks"] if b["t"] == "p" and b.get("lead"))
    assert p["num"] == 10 and '<span class="nl-lead">' in p["html"]
    assert '<mark class="nl-art">article L. 330-3</mark>' in p["html"]

    cps = detect_checkpoints(r["blocks"])
    assert len(cps) == 2
    assert all(r["blocks"][c["anchor"]]["t"] == "table" for c in cps)
    q = cps[0]["questions"]
    assert {x["type"] for x in q} == {"qcm"}  # 3 notions : que des QCM
    assert all(x["choices"][x["answer"]] in [n for n, _ in NOTIONS] for x in q)


def test_short_answers_tolerant():
    from app.services.lib_parser import check_short_answer

    assert check_short_answer("loi doubin", ["Loi Doubin"])
    assert check_short_answer("Doubin", ["Loi Doubin"])
    assert check_short_answer("Réseeau", ["Réseau"])  # faute de frappe
    assert not check_short_answer("loi", ["Loi Doubin"])
    assert check_short_answer("franchis", ["Franchise"])
    assert not check_short_answer("", ["Franchise"])
    assert not check_short_answer("concession", ["Franchise"])


def test_parcours_complet(client):
    # ── Console : import ──
    r = client.post(
        "/lib/admin/documents",
        headers=ADMIN,
        files={"file": ("Cours.docx", _cours_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"niveau": "M1", "matiere": "Distribution", "elo_palier": "10", "recompense_elo": "10"},
    )
    assert r.status_code == 200, r.text
    doc = r.json()
    assert doc["titre"] == "Droit de la distribution"
    assert doc["detection"]["points_passage"] == 2
    cp1, cp2 = doc["checkpoints"]
    assert (cp1["elo_requis"], cp2["elo_requis"]) == (0, 10)
    assert cp2["final"] is True and cp1["final"] is False   # après l'étape 2 : sources seulement

    assert client.post("/lib/admin/documents", files={"file": ("x.docx", b"x")}).status_code == 401

    # Brouillon : invisible pour l'élève
    uid, h = _user("lecteur")
    assert client.get(f"/lib/documents/{doc['id']}", headers=h).status_code == 404
    assert client.patch(f"/lib/admin/documents/{doc['id']}", headers=ADMIN, json={"is_published": True}).status_code == 200

    cat = client.get("/lib/catalogue", headers=h).json()
    item = next(i for i in cat["items"] if i["id"] == doc["id"])
    assert item["nb_points_passage"] == 2 and item["valides"] == 0 and item["nb_etapes"] == 2

    # ── Élève : lecture jusqu'au premier point de passage ──
    v = client.get(f"/lib/documents/{doc['id']}", headers=h).json()
    assert v["blocks"][-1]["id"] == f"b{cp1['anchor_block']}"
    assert v["teaser"] and v["gate"]["id"] == cp1["id"] and v["gate"]["status"] == "open"
    assert any(x["locked"] for x in v["plan"]) and not v["plan"][0]["locked"]

    # Le second point de passage n'est pas encore accessible
    assert client.get(f"/lib/checkpoints/{cp2['id']}/question", headers=h).status_code == 403

    # Une seule question par essai, et l'essai suivant en tire une autre
    q1 = client.get(f"/lib/checkpoints/{cp1['id']}/question", headers=h).json()["question"]
    r = _repondre(client, h, cp1["id"], juste=False)
    assert r["passed"] is False and r["elo_gagne"] == 0 and r["essais_restants"] == 2
    assert r["peut_passer"] is False and "expected" not in r["correction"]   # réponse cachée
    q2 = client.get(f"/lib/checkpoints/{cp1['id']}/question", headers=h).json()["question"]
    assert q2["id"] != q1["id"]
    # Répondre à une autre question que celle de l'essai : refusé
    assert client.post(f"/lib/checkpoints/{cp1['id']}/submit", headers=h,
                       json={"question_id": q1["id"], "answer": 0}).status_code == 409

    # Réussite : +10 Elo, la suite s'ouvre
    r = _repondre(client, h, cp1["id"])
    assert r["passed"] and r["correction"]["correct"] and r["elo_gagne"] == 10 and r["elo"] == 10
    assert r["unlocked"] is True

    # Revalider ne rapporte rien de plus
    r = _repondre(client, h, cp1["id"])
    assert r["deja_valide"] and r["elo_gagne"] == 0 and r["elo"] == 10

    # Point final : tout est lisible, le quiz final reste à faire
    v = client.get(f"/lib/documents/{doc['id']}", headers=h).json()
    assert len(v["blocks"]) == doc["nb_blocs"] and v["teaser"] == []
    assert v["gate"]["id"] == cp2["id"] and v["gate"]["final"] is True

    # ── Verrou d'Elo ──
    client.patch(f"/lib/admin/checkpoints/{cp2['id']}", headers=ADMIN, json={"elo_requis": 50})
    v = client.get(f"/lib/documents/{doc['id']}", headers=h).json()
    assert v["gate"]["status"] == "elo" and v["gate"]["elo"] == 10
    r = client.get(f"/lib/checkpoints/{cp2['id']}/question", headers=h)
    assert r.status_code == 403 and "50 Elo" in r.json()["detail"]

    # Elo du document entier
    client.patch(f"/lib/admin/documents/{doc['id']}", headers=ADMIN, json={"elo_requis": 500})
    v = client.get(f"/lib/documents/{doc['id']}", headers=h).json()
    assert v["gate"]["type"] == "document" and v["blocks"] == []
    client.patch(f"/lib/admin/documents/{doc['id']}", headers=ADMIN, json={"elo_requis": 0})

    # ── Console : statistiques, édition, nouvelle version ──
    a = client.get(f"/lib/admin/documents/{doc['id']}", headers=ADMIN).json()
    assert a["lecteurs"] == 1 and a["checkpoints"][0]["valides"] == 1 and a["blocks"]

    bad = client.patch(f"/lib/admin/checkpoints/{cp2['id']}", headers=ADMIN,
                       json={"questions": [{"id": "q1", "type": "qcm", "question": "Quoi ?", "choices": ["a"], "answer": 0}]})
    assert bad.status_code == 400
    ok = client.patch(f"/lib/admin/checkpoints/{cp2['id']}", headers=ADMIN, json={"questions": [
        {"id": "q1", "type": "qcm", "question": "Le DIP est remis…", "choices": ["20 jours avant", "Après"], "answer": 0, "points": 1},
        {"id": "q2", "type": "courte", "question": "Quelle loi ?", "accepted": ["Doubin"], "points": 3},
    ]})
    assert ok.status_code == 200 and ok.json()["source"] == "manuel"

    r = client.post(f"/lib/admin/documents/{doc['id']}/fichier", headers=ADMIN,
                    files={"file": ("Cours v2.docx", _cours_docx(), "application/octet-stream")})
    assert r.status_code == 200
    a = r.json()
    assert [c["id"] for c in a["checkpoints"]] == [cp1["id"], cp2["id"]]  # progression conservée
    assert a["checkpoints"][1]["source"] == "manuel"                       # questions éditées gardées
    assert a["checkpoints"][0]["valides"] == 1

    # ── Trois échecs : l'élève peut passer à la suite, sans Elo ──
    client.patch(f"/lib/admin/checkpoints/{cp2['id']}", headers=ADMIN, json={"elo_requis": 0})
    uid2, h2 = _user("tenace")
    assert client.post(f"/lib/checkpoints/{cp1['id']}/passer", headers=h2).status_code == 403
    for k in range(3):
        r = _repondre(client, h2, cp1["id"], juste=False)
        assert r["passed"] is False and r["peut_passer"] is (k == 2)
    assert r["correction"]["expected"]                                     # dévoilée au 3e échec
    v = client.get(f"/lib/documents/{doc['id']}", headers=h2).json()
    assert v["gate"]["id"] == cp1["id"] and v["gate"]["peut_passer"] is True
    r = client.post(f"/lib/checkpoints/{cp1['id']}/passer", headers=h2).json()
    assert r["passed"] and r["unlocked"]
    v = client.get(f"/lib/documents/{doc['id']}", headers=h2).json()
    assert v["checkpoints"][0]["skipped"] and v["gate"]["id"] == cp2["id"] and v["user"]["elo"] == 0

    assert client.get(f"/lib/admin/documents/{doc['id']}/source", headers=ADMIN).status_code == 200
    assert client.delete(f"/lib/admin/documents/{doc['id']}", headers=ADMIN).json()["ok"]
    assert client.get(f"/lib/documents/{doc['id']}", headers=h).status_code == 404


def test_lib_et_lib_plus(client, monkeypatch):
    r = client.post(
        "/lib/admin/documents", headers=ADMIN,
        files={"file": ("Cours.docx", _cours_docx(), "application/octet-stream")},
        data={"publier": "true"},
    )
    doc = r.json()
    plus = client.post(
        "/lib/admin/documents", headers=ADMIN,
        files={"file": ("Plus.docx", _cours_docx(), "application/octet-stream")},
        data={"publier": "true", "plus": "true"},
    ).json()
    assert plus["is_plus"] is True and doc["is_plus"] is False
    cp1 = doc["checkpoints"][0]

    _, gratuit = _user("gratuit", plan="free")
    cat = client.get("/lib/catalogue", headers=gratuit).json()
    assert cat["user"]["acces_plus"] is False
    assert next(i for i in cat["items"] if i["id"] == plus["id"])["plus"] is True

    # Lib : tout le document, sans point de passage, même avec un Elo requis
    client.patch(f"/lib/admin/documents/{doc['id']}", headers=ADMIN, json={"elo_requis": 500})
    v = client.get(f"/lib/documents/{doc['id']}?mode=lib", headers=gratuit).json()
    assert len(v["blocks"]) == doc["nb_blocs"] and v["checkpoints"] == [] and v["gate"] is None
    client.patch(f"/lib/admin/documents/{doc['id']}", headers=ADMIN, json={"elo_requis": 0})
    assert client.get(f"/lib/documents/{plus['id']}?mode=lib", headers=gratuit).status_code == 403

    # Lib+ : fermé au compte gratuit
    for url in (f"/lib/documents/{doc['id']}", f"/lib/checkpoints/{cp1['id']}/question"):
        r = client.get(url, headers=gratuit)
        assert r.status_code == 403 and "Lib+" in r.json()["detail"]
    assert client.post("/lib/chat", headers=gratuit, json={"document_id": doc["id"], "message": "?"}).status_code == 403

    # AdJuris actif : Lib+ ouvert
    from app.db.database import SessionLocal
    from app.db.models import PrepaAdjurisEnrollment
    uid, adj = _user("adjuris", plan="free")
    db = SessionLocal()
    db.add(PrepaAdjurisEnrollment(user_id=uid, matiere_key="l1_civil", status="active"))
    db.commit()
    db.close()
    assert client.get(f"/lib/documents/{plus['id']}", headers=adj).status_code == 200

    # Chat NAVIRE AI : nourri de l'extrait lu
    vu = {}

    def faux(**kw):
        vu.update(kw)
        return {"conversation_id": 7, "reply": "Réponse", "sources": []}

    import app.services.navire_chat as nc
    monkeypatch.setattr(nc, "send_message", faux)
    r = client.post("/lib/chat", headers=adj, json={"document_id": doc["id"], "checkpoint_id": cp1["id"], "message": "Et la franchise ?"})
    assert r.status_code == 200, r.text
    assert r.json() == {"conversation_id": 7, "reply": "Réponse", "sources": []}
    assert vu["user_message"] == "Et la franchise ?" and "Paragraphe 1-0" in vu["context"]
    assert "Paragraphe 2-0" not in vu["context"]                    # rien après le point de passage

    for d in (doc, plus):
        client.delete(f"/lib/admin/documents/{d['id']}", headers=ADMIN)


def test_pass_admin(client):
    from app.db.database import SessionLocal
    from app.db.models import User

    doc = client.post(
        "/lib/admin/documents", headers=ADMIN,
        files={"file": ("Cours.docx", _cours_docx(), "application/octet-stream")},
        data={"elo_palier": "50", "plus": "true"},          # brouillon, Lib+, Elo requis
    ).json()
    cp1, cp2 = doc["checkpoints"]
    client.patch(f"/lib/admin/documents/{doc['id']}", headers=ADMIN, json={"elo_requis": 500})

    uid, h = _user("admin", plan="free")
    db = SessionLocal()
    db.get(User, uid).is_admin = True
    db.commit()
    db.close()
    _, eleve = _user("eleve")

    # Brouillon : visible en aperçu par l'admin seulement
    cat = client.get("/lib/catalogue", headers=h).json()
    assert cat["user"]["admin"] and cat["user"]["acces_plus"]
    assert next(i for i in cat["items"] if i["id"] == doc["id"])["brouillon"] is True
    assert all(i["id"] != doc["id"] for i in client.get("/lib/catalogue", headers=eleve).json()["items"])
    assert client.get(f"/lib/documents/{doc['id']}", headers=eleve).status_code == 404

    # Aucun verrou d'Elo, et chaque point de passage se passe sans répondre
    v = client.get(f"/lib/documents/{doc['id']}", headers=h).json()
    assert v["gate"]["type"] == "checkpoint" and v["gate"]["peut_passer"] is True
    assert client.post(f"/lib/checkpoints/{cp1['id']}/passer", headers=h).json()["unlocked"] is True
    v = client.get(f"/lib/documents/{doc['id']}", headers=h).json()
    assert v["gate"]["id"] == cp2["id"] and v["gate"]["status"] == "open"   # 50 Elo requis : ignoré
    assert client.post(f"/lib/checkpoints/{cp2['id']}/passer", headers=h).json()["passed"]
    v = client.get(f"/lib/documents/{doc['id']}", headers=h).json()
    assert v["gate"] is None and len(v["blocks"]) == doc["nb_blocs"] and v["user"]["elo"] == 0

    client.delete(f"/lib/admin/documents/{doc['id']}", headers=ADMIN)


def _pdf_perso(paragraphes: int = 8) -> bytes:
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    y = 60
    phrase = "Le contrat de franchise suppose un savoir-faire secret, substantiel et identifié. "
    for k in range(paragraphes):
        for ligne in range(3):
            page.insert_text((60, y), f"{phrase[:80]} ({k}.{ligne})", fontname="helv", fontsize=9)
            y += 13
        y += 10
        if y > 760:
            page, y = pdf.new_page(), 60
    return pdf.tobytes()


def test_mes_fichiers(client, monkeypatch, tmp_path):
    import app.routers.lib as lib
    from app.db.database import SessionLocal
    from app.db.models import File

    monkeypatch.setattr(lib, "PERSO_MOTS", 60)
    uid, h = _user("perso", plan="membre+")
    _, autre = _user("autre", plan="membre+")
    _, gratuit_h = _user("gratuitperso", plan="free")

    chemin = tmp_path / "fiche.pdf"
    chemin.write_bytes(_pdf_perso())
    db = SessionLocal()
    f = File(user_id=uid, filename_original="Ma_fiche franchise.pdf", filename_stored="x.pdf",
             path=str(chemin), size_bytes=chemin.stat().st_size)
    db.add(f)
    db.commit()
    fid = f.id
    db.close()

    # Le fichier apparaît dans la bibliothèque de son propriétaire, en tête
    cat = client.get("/lib/catalogue", headers=h).json()
    perso = cat["items"][0]
    assert perso["perso"] is True and perso["kind"] == "perso" and perso["titre"] == "Ma fiche franchise"
    assert perso["plus"] is False and perso["brouillon"] is False
    assert all(i["id"] != perso["id"] for i in client.get("/lib/catalogue", headers=autre).json()["items"])
    assert client.get(f"/lib/documents/{perso['id']}?mode=lib", headers=autre).status_code == 404
    assert all(d["id"] != perso["id"] for d in client.get("/lib/admin/documents", headers=ADMIN).json()["items"])

    # Lib : lu à la première ouverture, en entier
    v = client.get(f"/lib/documents/{perso['id']}?mode=lib", headers=h).json()
    assert v["blocks"] and v["checkpoints"] == [] and v["gate"] is None

    # Lib+ : points de passage placés tout seuls, sans Elo
    v = client.get(f"/lib/documents/{perso['id']}", headers=h).json()
    assert len(v["checkpoints"]) >= 2 and v["gate"]["status"] == "open"
    assert all(c["elo_requis"] == 0 and c["recompense_elo"] == 0 for c in v["checkpoints"])
    cp1 = v["checkpoints"][0]

    # Questions rédigées par l'IA à l'arrivée sur le point de passage
    vu = {}

    def fausse_ia(texte, n_qcm=5, n_courte=2):
        vu["texte"] = texte
        return [{"id": "q1", "type": "qcm", "question": "Le savoir-faire doit être…", "choices": ["Secret", "Public"],
                 "answer": 0, "points": 1, "explication": ""}]

    monkeypatch.setattr(lib, "generer_questions_ia", fausse_ia)
    q = client.get(f"/lib/checkpoints/{cp1['id']}/question", headers=h).json()["question"]
    assert q["question"] == "Le savoir-faire doit être…" and "franchise" in vu["texte"]
    r = client.post(f"/lib/checkpoints/{cp1['id']}/submit", headers=h, json={"question_id": "q1", "answer": 0}).json()
    assert r["passed"] and r["elo_gagne"] == 0

    # Rien à demander : le point de passage s'efface, la lecture continue
    monkeypatch.setattr(lib, "generer_questions_ia", lambda *a, **k: [])
    cp2 = client.get(f"/lib/documents/{perso['id']}", headers=h).json()["gate"]
    r = client.get(f"/lib/checkpoints/{cp2['id']}/question", headers=h)
    assert r.status_code == 409
    v = client.get(f"/lib/documents/{perso['id']}", headers=h).json()
    assert all(c["id"] != cp2["id"] for c in v["checkpoints"])

    # Gratuit : ses fichiers dans Lib, pas de Lib+
    assert client.get("/lib/catalogue", headers=gratuit_h).status_code == 200

    # Fichier supprimé : le document disparaît
    db = SessionLocal()
    db.delete(db.get(File, fid))
    db.commit()
    db.close()
    assert all(i["id"] != perso["id"] for i in client.get("/lib/catalogue", headers=h).json()["items"])
    assert client.get(f"/lib/documents/{perso['id']}?mode=lib", headers=h).status_code == 404


def test_import_pdf(client):
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    y = 72
    lignes = [("PREMIÈRE PARTIE. LA FORMATION", True)] + [
        (f"{i}. Notion {i}.- Le contrat de franchise suppose un savoir-faire secret, substantiel et identifié, art. 1164 du Code civil.", False)
        for i in range(1, 9)
    ] + [("DEUXIÈME PARTIE. L'EXÉCUTION", True)] + [
        (f"{i}. Notion {i}.- La rupture brutale est sanctionnée par l'article L. 442-1 du Code de commerce.", False)
        for i in range(9, 13)
    ]
    for txt, bold in lignes:
        page.insert_text((72, y), txt[:95], fontname="hebo" if bold else "helv", fontsize=10)
        y += 24
    data = pdf.tobytes()

    from app.services.lib_parser import parse_document
    r = parse_document(data, "plan.pdf")
    assert [b["t"] for b in r["blocks"]].count("h1") == 2
    assert any(b.get("lead") == "Notion 3" for b in r["blocks"])

    r = client.post("/lib/admin/documents", headers=ADMIN,
                    files={"file": ("plan.pdf", data, "application/pdf")}, data={"kind": "fiche"})
    assert r.status_code == 200, r.text
    assert r.json()["kind"] == "fiche"
    client.delete(f"/lib/admin/documents/{r.json()['id']}", headers=ADMIN)

    r = client.post("/lib/admin/documents", headers=ADMIN, files={"file": ("x.txt", b"bonjour", "text/plain")})
    assert r.status_code == 400


def test_embed_servi(client):
    r = client.get("/lib/embed")
    assert r.status_code == 200 and 'API: ""' in r.text
